"""Excel/CSV import and export, including round-trip and injection protection (SYNTHETIC data)."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .api_helpers import generate, setup_team

SAMPLES = Path(__file__).resolve().parents[2] / "sample-data"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _preview(client, kind, name, data, **extra):
    r = client.post("/api/imports/preview", files={"file": (name, data, XLSX)}, data={"kind": kind, **extra})
    return r


def _commit(client, p, mapping=None, **options):
    return client.post("/api/imports/commit", json={"token": p["token"], "mapping": mapping or p["suggested_mapping"],
                                                    "options": options})


def test_export_has_all_sheets_and_formatting(client):
    p, analysts = setup_team(client)
    # A hostile name must come out as inert text.
    a = analysts[0]
    client.put(f"/api/analysts/{a['id']}", json={"employee_code": a["employee_code"],
               "full_name": "=HYPERLINK(\"http://evil\",\"x\")", "eligibility": [{"project_id": p["id"]}]})
    rid = generate(client, [p["id"]])["roster_id"]
    r = client.get(f"/api/rosters/{rid}/export.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith(XLSX)
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Monthly Roster", "Project Coverage", "Leave Register", "Workload Summary",
                             "Shift Legend", "Validation Report", "Configuration"]
    ws = wb["Monthly Roster"]
    assert ws.freeze_panes == "C5" and ws.auto_filter.ref.startswith("A4:")
    assert ws.cell(4, 1).value == "Employee Code"
    for row in ws.iter_rows(min_row=5, max_row=11):
        for c in row:
            assert not (isinstance(c.value, str) and c.value.startswith("=")), c.value
            assert c.data_type != "f"
    names = [ws.cell(r, 2).value for r in range(5, 12)]
    assert "'=HYPERLINK(\"http://evil\",\"x\")" in names
    filled = [c.fill.fgColor.rgb for c in ws[5][2:32] if c.value in ("M", "A", "N")]
    assert filled and all(f not in (None, "00000000") for f in filled)
    assert wb["Project Coverage"].tables and wb["Project Coverage"].freeze_panes == "A2"
    assert wb["Validation Report"].cell(2, 1).value.startswith("Result: VALID")
    cfg = {row[0].value: row[1].value for row in wb["Configuration"].iter_rows(min_row=2)}
    assert cfg["policy.min_rest_hours"] == 11
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:  # no formulas anywhere in the workbook XML
        assert not any(b"<f>" in z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/"))


def test_roster_export_import_round_trip(client):
    p, _ = setup_team(client)
    rid = generate(client, [p["id"]])["roster_id"]
    data = client.get(f"/api/rosters/{rid}/export.xlsx").content
    prev = _preview(client, "roster", "export.xlsx", data).json()
    assert prev["sheet"] == "Monthly Roster" and not prev["errors"]
    assert prev["detected"]["year"] == 2026 and prev["detected"]["month"] == 11
    assert prev["suggested_mapping"]["employee_code"] == "Employee Code"
    res = _commit(client, prev, year=2026, month=11, project_id=p["id"], status="draft").json()
    assert res["errors"] == [] and res["roster_id"], res
    key = lambda a: (a["analyst_id"], a["project_id"], a["shift_template_id"], a["date"])  # noqa: E731
    original = {key(a) for a in client.get(f"/api/rosters/{rid}").json()["assignments"]}
    imported = {key(a) for a in client.get(f"/api/rosters/{res['roster_id']}").json()["assignments"]}
    assert original == imported
    assert client.get(f"/api/rosters/{res['roster_id']}/validation").json()["is_valid"]
    # Re-export of the imported roster is identical in the roster grid.
    again = load_workbook(io.BytesIO(client.get(f"/api/rosters/{res['roster_id']}/export.xlsx").content))
    first = load_workbook(io.BytesIO(data))
    grid = lambda wb: [[c.value for c in r] for r in wb["Monthly Roster"].iter_rows(min_row=4, max_row=11, max_col=32)]  # noqa: E731
    assert grid(again) == grid(first)


def test_historical_synthetic_roster_adapter(client):
    p = client.post("/api/projects", json={"code": "SOC1", "name": "Synthetic SOC"}).json()
    data = (SAMPLES / "SYNTHETIC_historical_soc_roster_2026-08.xlsx").read_bytes()
    prev = _preview(client, "roster", "hist.xlsx", data).json()
    assert prev["suggested_mapping"]["employee_code"] == "Emp Code"
    assert {"M", "A", "N", "WO", "L"} <= set(prev["detected"]["codes_found"])
    r = _commit(client, prev, year=2026, month=8, project_id=p["id"])
    assert r.json()["errors"] and "not found" in r.json()["errors"][0]["message"]  # all-or-nothing
    assert client.get("/api/rosters").json() == []
    prev = _preview(client, "roster", "hist.xlsx", data).json()
    res = _commit(client, prev, year=2026, month=8, project_id=p["id"], create_missing_analysts=True).json()
    assert res["roster_id"] and "created 7 analyst" in res["message"]
    roster = client.get(f"/api/rosters/{res['roster_id']}").json()
    assert roster["status"] == "historical" and len(roster["rows"]) == 7
    markers = {c["marker"] for r in roster["rows"] for c in r["cells"].values()}
    assert {"WO", "LEAVE"} <= markers
    # The historical roster feeds month continuity for September.
    job = generate(client, [p["id"]], year=2026, month=9)
    assert any("Continuity from previous month" in n for n in job["result"]["notes"]), job["result"]["notes"]


def test_unknown_code_and_formula_rejected(client):
    p = client.post("/api/projects", json={"code": "SOC1", "name": "Synthetic"}).json()
    header = ["Emp Code", "Name", *range(1, 31)]
    row = ["SYN-1", "Synthetic", "M", "ZZ", *["WO"] * 28]
    prev = _preview(client, "roster", "r.xlsx", _xlsx([header, row])).json()
    res = _commit(client, prev, year=2026, month=11, project_id=p["id"], create_missing_analysts=True).json()
    assert any("unknown code 'ZZ'" in e["message"] for e in res["errors"])
    row = ["SYN-1", "Synthetic", "=1+1", *["WO"] * 29]
    prev = _preview(client, "roster", "r.xlsx", _xlsx([header, row])).json()
    assert any("formulas" in e for e in prev["errors"])


def test_analyst_import_xlsx_and_csv(client):
    client.post("/api/projects", json={"code": "SOC1", "name": "Synthetic"})
    data = (SAMPLES / "SYNTHETIC_analysts.csv").read_bytes()
    prev = client.post("/api/imports/preview", files={"file": ("a.csv", data, "text/csv")},
                       data={"kind": "analysts"}).json()
    assert prev["suggested_mapping"]["employee_code"] == "Employee Code"
    assert prev["row_count"] == 7
    res = _commit(client, prev).json()
    assert res["created"] == 7, res
    analysts = client.get("/api/analysts").json()
    assert all(a["eligibility"][0]["is_primary"] for a in analysts)
    skills = client.get("/api/skills").json()
    assert [s["name"] for s in skills] == ["SIEM"]
    # Same file again: skipped, or updated when asked.
    prev = client.post("/api/imports/preview", files={"file": ("a.csv", data, "text/csv")}, data={"kind": "analysts"}).json()
    assert _commit(client, prev).json()["skipped"] == 7
    bad = _xlsx([["Employee Code", "Full Name", "Projects"], ["SYN-X", "Ok", "SOC1"], ["SYN-Y", "", "NOPE"],
                 ["SYN-X", "Dup", "SOC1"]])
    prev = _preview(client, "analysts", "b.xlsx", bad).json()
    res = _commit(client, prev).json()
    assert res["created"] == 0 and len(res["errors"]) >= 2
    assert not [a for a in client.get("/api/analysts").json() if a["employee_code"] == "SYN-X"]


def test_leave_import(client):
    p, _ = setup_team(client, code="SYN")
    rows = [["Employee Code", "Start Date", "End Date", "Leave Type", "Status", "Start Time", "End Time"],
            ["SYN-002", "2026-11-09", "2026-11-11", "annual", "approved", None, None],
            ["SYN-004", "18/11/2026", None, "casual", "pending", None, None],
            ["SYN-006", "2026-11-05", "2026-11-05", "sick", "approved", "09:00", "13:00"]]
    prev = _preview(client, "leave", "l.xlsx", _xlsx(rows)).json()
    res = _commit(client, prev).json()
    assert res["created"] == 3, res
    leaves = client.get("/api/leaves").json()
    assert {lv["status"] for lv in leaves} == {"approved", "pending"}
    assert any(lv["partial_day"] and lv["start_time"] == "09:00" for lv in leaves)
    assert all(lv["source"] == "import" for lv in leaves)
    prev = _preview(client, "leave", "l.xlsx", _xlsx(rows)).json()
    assert _commit(client, prev).json()["skipped"] == 2  # exact duplicates are not re-added
    bad = [rows[0], ["SYN-999", "2026-11-01", None, "annual", "approved", None, None],
           ["SYN-001", "not a date", None, "annual", "approved", None, None],
           ["SYN-002", "2026-11-10", "2026-11-14", "annual", "approved", None, None]]
    prev = _preview(client, "leave", "b.xlsx", _xlsx(bad)).json()
    res = _commit(client, prev).json()
    msgs = " ".join(e["message"] for e in res["errors"])
    assert res["created"] == 0 and "unknown employee code" in msgs and "not a date" in msgs and "overlaps" in msgs


def test_upload_restrictions(client, settings):
    assert _preview(client, "analysts", "a.xls", b"whatever").status_code == 415
    assert _preview(client, "analysts", "a.xlsx", b"not a zip").status_code == 422
    assert _preview(client, "analysts", "a.xlsx", b"").status_code == 422
    assert _preview(client, "bogus", "a.csv", b"a,b\n1,2\n").status_code == 422
    big = b"x" * (settings.max_upload_bytes + 10)
    assert _preview(client, "analysts", "a.csv", big).status_code == 413
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/vbaProject.bin", b"macro")
    assert _preview(client, "analysts", "m.xlsx", buf.getvalue()).status_code == 422
    r = client.post("/api/imports/commit", json={"token": "nope", "mapping": {}, "options": {}})
    assert r.status_code == 410


def test_import_templates(client):
    r = client.get("/api/imports/templates/analysts.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.active.cell(1, 1).value == "Employee Code"
    assert client.get("/api/imports/templates/nope.xlsx").status_code == 404
