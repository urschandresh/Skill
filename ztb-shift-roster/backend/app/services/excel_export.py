"""Editable XLSX export of a roster (7 sheets)."""

from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import APP_NAME, APP_VERSION
from ..models import Analyst, LeaveRequest, Roster, ShiftCodeMapping
from . import roster_service as svc
from .excel_common import (
    BORDER,
    CENTER,
    HEADER_FILL,
    HEADER_FONT,
    MARKER_FILLS,
    TITLE_FONT,
    add_table,
    fill_for,
    write_cell,
)
from .problem_builder import month_bounds

ROSTER_HEADER_ROW = 4
FIXED_COLUMNS = ["Employee Code", "Analyst"]


def cell_code(assignments: list[dict], templates: dict[int, dict], projects: dict[int, str], multi: bool) -> str:
    parts = []
    for a in sorted(assignments, key=lambda x: templates[x["shift_template_id"]]["start_time"]):
        code = templates[a["shift_template_id"]]["code"]
        parts.append(f"{projects[a['project_id']]}:{code}" if multi else code)
    return " / ".join(parts)


def export_roster(session: Session, roster: Roster) -> bytes:
    detail = svc.roster_detail(session, roster)
    inst, report = svc.validate_stored(session, roster)
    validation = report.to_dict()
    templates = {t["id"]: t for t in detail["shift_templates"]}
    projects = {p["id"]: p["code"] for p in detail["projects"]}
    multi = len(roster.project_ids or []) > 1 or len({a["project_id"] for a in detail["assignments"]}) > 1
    start, end = month_bounds(roster.year, roster.month)
    dates = [date.fromisoformat(d) for d in detail["rows"][0]["cells"]] if detail["rows"] else []

    wb = Workbook()
    # ------------------------------------------------------------ Monthly Roster
    ws = wb.active
    ws.title = "Monthly Roster"
    write_cell(ws, 1, 1, f"{APP_NAME} · {roster.name}", font=TITLE_FONT)
    write_cell(ws, 2, 1, f"Status: {roster.status} · version {roster.version} · "
                         f"projects {', '.join(projects[p] for p in roster.project_ids or [] if p in projects)} · "
                         f"exported {datetime.now():%Y-%m-%d %H:%M}")
    headers = FIXED_COLUMNS + dates + ["Shifts", "Nights", "Hours"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=ROSTER_HEADER_ROW, column=c, value=h if not isinstance(h, str) else None)
        if isinstance(h, str):
            write_cell(ws, ROSTER_HEADER_ROW, c, h)
        else:
            cell.number_format = "dd ddd"
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, CENTER, BORDER
    metrics = validation["metrics"]
    for r, row in enumerate(detail["rows"], start=ROSTER_HEADER_ROW + 1):
        write_cell(ws, r, 1, row["employee_code"], border=BORDER)
        write_cell(ws, r, 2, row["full_name"], border=BORDER)
        for c, d in enumerate(dates, start=3):
            cell_data = row["cells"][d.isoformat()]
            if cell_data["assignments"]:
                value = cell_code(cell_data["assignments"], templates, projects, multi)
                fill = fill_for(templates[cell_data["assignments"][0]["shift_template_id"]]["color"])
            else:
                value = cell_data["marker"] or "WO"
                fill = MARKER_FILLS.get(value)
            cell = write_cell(ws, r, c, value, alignment=CENTER, border=BORDER)
            if fill is not None:
                cell.fill = fill
            if any(a["locked"] for a in cell_data["assignments"]):
                cell.font = Font(bold=True)
        m = metrics.get(str(row["analyst_id"]), {})
        for k, key in enumerate(("shifts", "nights", "hours")):
            write_cell(ws, r, 3 + len(dates) + k, m.get(key, 0), border=BORDER)
    last_row = ROSTER_HEADER_ROW + max(1, len(detail["rows"]))
    ws.freeze_panes = ws.cell(row=ROSTER_HEADER_ROW + 1, column=3)
    ws.auto_filter.ref = f"A{ROSTER_HEADER_ROW}:{get_column_letter(len(headers))}{last_row}"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 26
    for c in range(3, 3 + len(dates)):
        ws.column_dimensions[get_column_letter(c)].width = 11 if multi else 7
    write_cell(ws, last_row + 2, 1, "Bold = locked assignment. Edit codes using the Shift Legend sheet; "
                                     "re-import via Import / Export > Roster.", font=Font(italic=True, size=9))

    # ------------------------------------------------------------ Project Coverage
    ws = wb.create_sheet("Project Coverage")
    names = {r["analyst_id"]: r["employee_code"] for r in detail["rows"]}
    rows = []
    for c in report.coverage:
        who = sorted(names.get(a["analyst_id"], str(a["analyst_id"])) for a in detail["assignments"]
                     if a["shift_template_id"] == c["shift_id"] and a["date"] == c["date"])
        status = "SHORT" if c["assigned"] < c["required"] else ("OVER MAX" if c["max"] is not None
                                                               and c["assigned"] > c["max"] else "OK")
        rows.append([date.fromisoformat(c["date"]), projects.get(c["project_id"], c["project_id"]),
                     templates.get(c["shift_id"], {}).get("code", c["shift_id"]), c["required"],
                     c["max"] if c["max"] is not None else "", c["assigned"], status, ", ".join(who)])
    add_table(ws, "Coverage", ["Date", "Project", "Shift", "Required", "Max", "Assigned", "Status", "Analysts"], rows,
              widths={"Analysts": 40, "Date": 12})
    for r in range(2, len(rows) + 2):
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"

    # ------------------------------------------------------------ Leave Register
    ws = wb.create_sheet("Leave Register")
    ids = [r["analyst_id"] for r in detail["rows"]]
    analysts = {a.id: a for a in session.scalars(select(Analyst).where(Analyst.id.in_(ids)))}
    leaves = session.scalars(select(LeaveRequest).where(
        LeaveRequest.analyst_id.in_(ids), LeaveRequest.start_date <= end, LeaveRequest.end_date >= start)
        .order_by(LeaveRequest.start_date))
    rows = [[analysts[lv.analyst_id].employee_code, analysts[lv.analyst_id].full_name, lv.start_date, lv.end_date,
             lv.leave_type, "yes" if lv.partial_day else "no", lv.start_time or "", lv.end_time or "", lv.status,
             lv.reason] for lv in leaves]
    add_table(ws, "Leave", ["Employee Code", "Analyst", "Start", "End", "Type", "Partial", "From", "To", "Status",
                            "Reason"], rows, widths={"Start": 12, "End": 12, "Reason": 40})
    for r in range(2, len(rows) + 2):
        for c in (3, 4):
            ws.cell(row=r, column=c).number_format = "yyyy-mm-dd"

    # ------------------------------------------------------------ Workload Summary
    ws = wb.create_sheet("Workload Summary")
    codes = sorted({t["code"] for t in templates.values()})
    rows = []
    for row in detail["rows"]:
        m = metrics.get(str(row["analyst_id"]), {})
        rows.append([row["employee_code"], row["full_name"], m.get("shifts", 0), m.get("worked_days", 0),
                     m.get("hours", 0), m.get("nights", 0), m.get("weekend_shifts", 0), m.get("leave_days", 0),
                     *[m.get("by_code", {}).get(c, 0) for c in codes]])
    add_table(ws, "Workload", ["Employee Code", "Analyst", "Shifts", "Worked days", "Hours", "Nights",
                               "Weekend shifts", "Leave days", *[f"{c} shifts" for c in codes]], rows)

    # ------------------------------------------------------------ Shift Legend
    ws = wb.create_sheet("Shift Legend")
    rows = [[projects.get(t["project_id"], t["project_id"]), t["code"], t["name"], t["start_time"], t["end_time"],
             "yes" if t["end_time"] <= t["start_time"] else "no", "yes" if t["is_night"] else "no", t["color"]]
            for t in sorted(templates.values(), key=lambda t: (t["project_id"], t["start_time"]))]
    rows += [["", "WO", "Weekly off", "", "", "", "", ""], ["", "LEAVE", "Approved leave", "", "", "", "", ""],
             ["", "H", "Holiday (work-blocking)", "", "", "", "", ""]]
    add_table(ws, "Legend", ["Project", "Code", "Name", "Start", "End", "Ends next day", "Night", "Colour"], rows)
    for r, t in enumerate(sorted(templates.values(), key=lambda t: (t["project_id"], t["start_time"])), start=2):
        ws.cell(row=r, column=2).fill = fill_for(t["color"])
    if multi:
        write_cell(ws, len(rows) + 3, 1, "In the roster sheet, codes are written PROJECT:CODE because this roster "
                                          "covers several projects.", font=Font(italic=True))
    mrow = len(rows) + 5
    write_cell(ws, mrow, 1, "Historical code mappings used for imports", font=Font(bold=True))
    for i, m in enumerate(session.scalars(select(ShiftCodeMapping).order_by(ShiftCodeMapping.source_code)), start=1):
        write_cell(ws, mrow + i, 1, m.source_code)
        write_cell(ws, mrow + i, 2, m.target_kind if m.target_kind != "shift" else f"shift {m.target_shift_code}")
        write_cell(ws, mrow + i, 3, m.description)

    # ------------------------------------------------------------ Validation Report
    ws = wb.create_sheet("Validation Report")
    write_cell(ws, 1, 1, "Independent validation", font=TITLE_FONT)
    write_cell(ws, 2, 1, f"Result: {'VALID' if validation['is_valid'] else 'VIOLATIONS FOUND'} · "
                         f"{validation['error_count']} error(s), {validation['warning_count']} warning(s) · "
                         f"coverage {validation['coverage_percent']}%")
    rows = [[v["severity"], v["rule"], v["date"] or "", names.get(v["analyst_id"], "") if v["analyst_id"] else "",
             projects.get(v["project_id"], "") if v["project_id"] else "", v["message"]]
            for v in validation["violations"]]
    add_table(ws, "Violations", ["Severity", "Rule", "Date", "Analyst", "Project", "Message"], rows, start_row=4,
              widths={"Message": 90})

    # ------------------------------------------------------------ Configuration
    ws = wb.create_sheet("Configuration")
    snap = dict(roster.policy_snapshot or {})
    weights = snap.pop("weights", {}) or {}
    rows = [["Application", f"{APP_NAME} {APP_VERSION}"], ["Roster", roster.name], ["Roster id", roster.id],
            ["Month", f"{roster.year}-{roster.month:02d}"], ["Status", roster.status], ["Version", roster.version],
            ["Solver status", roster.solver_status or ""],
            ["Objective", (roster.solver_summary or {}).get("objective", "")],
            ["Best bound", (roster.solver_summary or {}).get("best_bound", "")],
            ["Projects", ", ".join(projects.get(p, str(p)) for p in roster.project_ids or [])],
            ["Exported at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")]]
    rows += [[f"policy.{k}", v if not isinstance(v, (dict, list)) else str(v)] for k, v in sorted(snap.items())]
    rows += [[f"weight.{k}", v] for k, v in sorted(weights.items())]
    rows += [["notes", roster.notes or ""]]
    add_table(ws, "Configuration", ["Setting", "Value"], rows, widths={"Setting": 34, "Value": 50})
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(horizontal="left")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


TEMPLATES = {
    "analysts": (["Employee Code", "Full Name", "Email", "Projects", "Primary Project", "Skills", "Active",
                  "Max Hours Per Week"],
                 [["SYN-101", "Example Analyst (synthetic)", "", "SOC1, SOC2", "SOC1", "SIEM", "yes", ""]]),
    "leave": (["Employee Code", "Start Date", "End Date", "Leave Type", "Status", "Start Time", "End Time", "Reason"],
              [["SYN-101", "2026-11-10", "2026-11-12", "annual", "approved", "", "", "Example (synthetic)"]]),
}


def import_template(kind: str) -> bytes:
    headers, rows = TEMPLATES[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = kind.capitalize()
    add_table(ws, kind.capitalize(), headers, rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
