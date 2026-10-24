"""Spreadsheet helpers shared by import and export: formula-injection guards and styling."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# A cell starting with one of these may be evaluated as a formula by Excel/LibreOffice
# (or by a CSV consumer). OWASP: prefix such text with a single quote.
DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")

NAVY = "0A1628"
HEADER_FILL = PatternFill("solid", fgColor="1F3A68")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
THIN = Side(style="thin", color="C7D2E3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

MARKER_FILLS = {
    "WO": PatternFill("solid", fgColor="E5E9F0"),
    "LEAVE": PatternFill("solid", fgColor="FDE68A"),
    "H": PatternFill("solid", fgColor="DDD6FE"),
}


def safe_text(value):
    """Neutralise text that a spreadsheet could execute. Non-strings pass through unchanged."""
    if isinstance(value, str) and value.startswith(DANGEROUS_PREFIXES):
        return "'" + value
    return value


def unescape_text(value):
    """Reverse ``safe_text`` for values read back from our own exports."""
    if isinstance(value, str) and len(value) > 1 and value[0] == "'" and value[1:].startswith(DANGEROUS_PREFIXES):
        return value[1:]
    return value


def write_cell(ws, row: int, col: int, value, **style):
    cell = ws.cell(row=row, column=col, value=safe_text(value))
    if isinstance(value, str):
        cell.data_type = "s"  # never let openpyxl store a string as a formula
    for k, v in style.items():
        setattr(cell, k, v)
    return cell


def fill_for(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color.lstrip("#").upper())


def add_table(ws, name: str, headers: list[str], rows: list[list], start_row: int = 1, widths: dict | None = None):
    """Write a formatted Excel table (with filter buttons) and freeze its header row."""
    for c, h in enumerate(headers, start=1):
        write_cell(ws, start_row, c, h, font=HEADER_FONT, fill=HEADER_FILL, alignment=CENTER, border=BORDER)
    for r, row in enumerate(rows, start=start_row + 1):
        for c, v in enumerate(row, start=1):
            write_cell(ws, r, c, v, border=BORDER)
    last_row = start_row + max(len(rows), 1)
    if not rows:  # tables need at least one data row
        for c in range(1, len(headers) + 1):
            write_cell(ws, last_row, c, "", border=BORDER)
    ref = f"A{start_row}:{get_column_letter(len(headers))}{last_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    for c, h in enumerate(headers, start=1):
        width = (widths or {}).get(h) or min(45, max(10, len(str(h)) + 2,
                                                      *(len(str(r[c - 1])) + 2 for r in rows[:200] if c - 1 < len(r))))
        ws.column_dimensions[get_column_letter(c)].width = width
