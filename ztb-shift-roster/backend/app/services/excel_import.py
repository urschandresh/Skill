"""XLSX/CSV import: safe reading, preview with column mapping, and all-or-nothing commits.

Imports are validated completely before anything is written. If any row has a problem the
whole import is rejected and every problem is listed, so a file is never half-imported.
"""

from __future__ import annotations

import io
import re
import threading
import time
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

import pandas as pd
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import (
    Analyst,
    AnalystProject,
    Assignment,
    LeaveRequest,
    Project,
    Roster,
    RosterDayMarker,
    ShiftCodeMapping,
    ShiftTemplate,
    Skill,
)
from ..schemas.organization import AnalystIn
from ..schemas.scheduling import LeaveIn
from .audit import record
from .excel_common import unescape_text
from .problem_builder import month_bounds
from .roster_service import next_version

MAX_ROWS = 5000
MAX_COLS = 200
MAX_UNCOMPRESSED = 60 * 1024 * 1024
PREVIEW_ROWS = 20
TOKEN_TTL = 30 * 60


class FormulaCell(str):
    """Marks a cell that contained a formula; never imported."""


@dataclass
class Sheet:
    name: str
    rows: list[list]


@dataclass
class Upload:
    kind: str
    filename: str
    sheets: list[Sheet]
    sheet: str
    created: float = field(default_factory=time.monotonic)


class ImportStore:
    """Parsed uploads awaiting commit, held in memory only (never written to disk)."""

    def __init__(self) -> None:
        self._items: dict[str, Upload] = {}
        self._lock = threading.Lock()

    def put(self, upload: Upload) -> str:
        token = str(uuid.uuid4())
        with self._lock:
            now = time.monotonic()
            for k in [k for k, v in self._items.items() if now - v.created > TOKEN_TTL]:
                del self._items[k]
            while len(self._items) >= 20:
                self._items.pop(next(iter(self._items)))
            self._items[token] = upload
        return token

    def get(self, token: str) -> Upload:
        with self._lock:
            up = self._items.get(token)
        if up is None or time.monotonic() - up.created > TOKEN_TTL:
            raise HTTPException(status_code=410, detail="Preview expired; upload the file again")
        return up

    def drop(self, token: str) -> None:
        with self._lock:
            self._items.pop(token, None)


# --------------------------------------------------------------------------- reading


def _clean(v):
    if isinstance(v, str):
        v = unescape_text(v.strip())
        return v if v != "" else None
    return v


def read_file(filename: str, data: bytes) -> list[Sheet]:
    name = filename.lower()
    if name.endswith(".csv"):
        return [_read_csv(data)]
    if name.endswith(".xlsx"):
        return _read_xlsx(data)
    raise HTTPException(status_code=415, detail="Only .xlsx and .csv files are accepted")


def _read_csv(data: bytes) -> Sheet:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 decodes almost anything
        raise HTTPException(status_code=422, detail="Could not decode the CSV file")
    if "\x00" in text:
        raise HTTPException(status_code=422, detail="CSV contains binary data")
    try:
        df = pd.read_csv(io.StringIO(text), dtype=str, header=None, keep_default_na=False, nrows=MAX_ROWS + 1,
                         sep=None, engine="python")
    except Exception as exc:  # pandas raises several parser error types
        raise HTTPException(status_code=422, detail="Could not parse the CSV file") from exc
    if df.shape[0] > MAX_ROWS or df.shape[1] > MAX_COLS:
        raise HTTPException(status_code=413, detail=f"File too large (max {MAX_ROWS} rows, {MAX_COLS} columns)")
    rows = [[_csv_value(v) for v in r] for r in df.itertuples(index=False, name=None)]
    return Sheet("CSV", rows)


def _csv_value(v: str):
    v = _clean(v)
    if isinstance(v, str) and v.startswith("="):
        return FormulaCell(v)
    return v


def _read_xlsx(data: bytes) -> list[Sheet]:
    if not data.startswith(b"PK\x03\x04"):
        raise HTTPException(status_code=422, detail="Not a valid .xlsx file")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
                raise HTTPException(status_code=413, detail="Workbook expands to too much data")
            if any(i.filename.lower().endswith("vbaproject.bin") for i in infos):
                raise HTTPException(status_code=422, detail="Macro-enabled workbooks are not accepted")
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=422, detail="Not a valid .xlsx file") from exc
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Could not open the workbook") from exc
    sheets = []
    try:
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(max_col=MAX_COLS)):
                if i >= MAX_ROWS:
                    raise HTTPException(status_code=413, detail=f"Sheet '{ws.title}' has more than {MAX_ROWS} rows")
                values = []
                for cell in row:
                    v = cell.value
                    if getattr(cell, "data_type", None) == "f" or (isinstance(v, str) and v.startswith("=")):
                        values.append(FormulaCell(str(v)))
                    else:
                        values.append(_clean(v))
                rows.append(values)
            while rows and all(v is None for v in rows[-1]):
                rows.pop()
            sheets.append(Sheet(ws.title, rows))
    finally:
        wb.close()
    if not sheets:
        raise HTTPException(status_code=422, detail="Workbook has no sheets")
    return sheets


# --------------------------------------------------------------------------- fields & mapping

FIELDS = {
    "analysts": [
        ("employee_code", True, "Unique employee code", ["employeecode", "empcode", "code", "employeeid", "empid", "id"]),
        ("full_name", True, "Full name", ["fullname", "name", "analyst", "analystname", "employeename"]),
        ("email", False, "Email (optional)", ["email", "mail", "emailaddress"]),
        ("projects", False, "Eligible project codes, comma separated", ["projects", "project", "projectcodes"]),
        ("primary_project", False, "Primary project code", ["primaryproject", "primary"]),
        ("skills", False, "Skill names, comma separated (created if new)", ["skills", "skill"]),
        ("active", False, "yes/no (default yes)", ["active", "status", "enabled"]),
        ("max_hours_per_week", False, "Whole hours (blank = policy)", ["maxhoursperweek", "maxhours", "weeklyhours"]),
    ],
    "leave": [
        ("employee_code", True, "Employee code of the analyst", ["employeecode", "empcode", "code", "employeeid", "id"]),
        ("start_date", True, "First day of leave", ["startdate", "start", "from", "fromdate", "leavefrom"]),
        ("end_date", False, "Last day (blank = start date)", ["enddate", "end", "to", "todate", "leaveto"]),
        ("leave_type", False, "annual/sick/casual/comp_off/unpaid/other", ["leavetype", "type"]),
        ("status", False, "approved or pending (default pending)", ["status", "approval"]),
        ("start_time", False, "Partial day start HH:MM", ["starttime", "fromtime"]),
        ("end_time", False, "Partial day end HH:MM", ["endtime", "totime"]),
        ("reason", False, "Reason", ["reason", "remarks", "comment", "comments"]),
    ],
    "roster": [
        ("employee_code", False, "Analyst code column (code or name required)", ["employeecode", "empcode", "code",
                                                                                 "employeeid", "empid", "id"]),
        ("full_name", False, "Analyst name column", ["analyst", "name", "fullname", "analystname", "employeename"]),
    ],
}


def _norm(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h).lower()) if h is not None else ""


def _date_header(v, year: int | None = None, month: int | None = None) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and float(v).is_integer() and 1 <= int(v) <= 31:
        return date(year, month, int(v)) if year and month and int(v) <= month_bounds(year, month)[1].day else None
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d-%b-%y", "%d %b %Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
        m = re.match(r"^(\d{1,2})(?:\s*[A-Za-z]{2,3})?$", s)  # "1", "01 Sun"
        if m and year and month and 1 <= int(m.group(1)) <= month_bounds(year, month)[1].day:
            return date(year, month, int(m.group(1)))
    return None


def _is_day_header(v) -> bool:
    if isinstance(v, (datetime, date)):
        return True
    if isinstance(v, (int, float)) and float(v).is_integer() and 1 <= int(v) <= 31:
        return True
    return isinstance(v, str) and (_date_header(v) is not None or re.match(r"^\d{1,2}(\s*[A-Za-z]{2,3})?$", v.strip()))


def find_header(rows: list[list], kind: str) -> int:
    if kind == "roster":
        for i, row in enumerate(rows[:20]):
            if sum(1 for v in row if v is not None and _is_day_header(v)) >= 7:
                return i
        raise HTTPException(status_code=422, detail="Could not find a header row with day columns (1..31 or dates)")
    for i, row in enumerate(rows[:20]):
        if sum(1 for v in row if v is not None) >= 2:
            return i
    raise HTTPException(status_code=422, detail="Could not find a header row")


def _headers(row: list) -> list[str]:
    out, seen = [], set()
    for i, v in enumerate(row):
        h = str(v.date() if isinstance(v, datetime) else v) if v is not None else f"Column {i + 1}"
        while h in seen:
            h += "_"
        seen.add(h)
        out.append(h)
    return out


def preview(store: ImportStore, kind: str, filename: str, data: bytes, sheet: str | None) -> dict:
    if kind not in FIELDS:
        raise HTTPException(status_code=422, detail="kind must be analysts, leave or roster")
    sheets = read_file(filename, data)
    chosen = next((s for s in sheets if s.name == sheet), None) if sheet else None
    if chosen is None:
        chosen = sheets[0]
        if kind == "roster":  # prefer the sheet that looks like a roster
            for s in sheets:
                try:
                    find_header(s.rows, "roster")
                    chosen = s
                    break
                except HTTPException:
                    continue
    errors, warnings, detected = [], [], {}
    columns, body = [], []
    try:
        h = find_header(chosen.rows, kind)
        columns = _headers(chosen.rows[h])
        body = chosen.rows[h + 1:]
    except HTTPException as exc:
        errors.append(str(exc.detail))
    norm = {_norm(c): c for c in columns}
    mapping = {}
    for name, _req, _desc, synonyms in FIELDS[kind]:
        mapping[name] = next((norm[s] for s in synonyms if s in norm), None)
    if kind == "roster" and columns:
        hdr = chosen.rows[h]
        days = [_date_header(v) for v in hdr if v is not None and _is_day_header(v)]
        real = [d for d in days if d is not None]
        detected["day_columns"] = sum(1 for v in hdr if v is not None and _is_day_header(v))
        if real:
            detected["year"], detected["month"] = real[0].year, real[0].month
        codes = sorted({str(v).upper() for r in body for i, v in enumerate(r)
                        if i < len(hdr) and hdr[i] is not None and _is_day_header(hdr[i]) and isinstance(v, str)})
        detected["codes_found"] = codes[:40]
        if not mapping["employee_code"] and not mapping["full_name"]:
            mapping["full_name"] = columns[0]
            warnings.append(f"Assuming column '{columns[0]}' identifies the analyst")
    formulas = sum(1 for r in body for v in r if isinstance(v, FormulaCell))
    if formulas:
        errors.append(f"{formulas} cell(s) contain formulas; formulas are not accepted. Save values only.")
    token = store.put(Upload(kind=kind, filename=filename, sheets=sheets, sheet=chosen.name))
    rows_out = [[None if v is None else str(v.date() if isinstance(v, datetime) else v) for v in r[:len(columns)]]
                for r in body if any(v is not None for v in r)]
    return {
        "token": token, "kind": kind, "filename": filename, "sheets": [s.name for s in sheets],
        "sheet": chosen.name, "columns": columns, "rows": rows_out[:PREVIEW_ROWS], "row_count": len(rows_out),
        "fields": [{"name": n, "required": r, "description": d} for n, r, d, _ in FIELDS[kind]],
        "suggested_mapping": mapping, "detected": detected, "errors": errors, "warnings": warnings,
    }


# --------------------------------------------------------------------------- commit


class RowErrors:
    def __init__(self) -> None:
        self.items: list[dict] = []

    def add(self, row: int | None, message: str) -> None:
        if len(self.items) < 500:
            self.items.append({"row": row, "message": message})

    def __bool__(self) -> bool:
        return bool(self.items)


def _records(upload: Upload, mapping: dict[str, str | None]):
    sheet = next(s for s in upload.sheets if s.name == upload.sheet)
    h = find_header(sheet.rows, upload.kind)
    columns = _headers(sheet.rows[h])
    index = {c: i for i, c in enumerate(columns)}
    for name, col in mapping.items():
        if col is not None and col not in index:
            raise HTTPException(status_code=422, detail=f"Mapped column '{col}' for {name} is not in the file")
    for offset, row in enumerate(sheet.rows[h + 1:], start=h + 2):
        if not any(v is not None for v in row):
            continue
        rec = {}
        for name, col in mapping.items():
            v = row[index[col]] if col is not None and index[col] < len(row) else None
            rec[name] = v
        yield offset, rec, row, columns


def _text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, FormulaCell):
        raise ValueError("formulas are not accepted")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        v = v.date()
    return str(v).strip() or None


def _parse_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _text(v)
    if s:
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    raise ValueError(f"'{v}' is not a date (use YYYY-MM-DD)")


def _parse_time(v) -> str | None:
    if v is None:
        return None
    if hasattr(v, "strftime") and not isinstance(v, (date, datetime)):
        return v.strftime("%H:%M")
    if isinstance(v, datetime):
        return v.strftime("%H:%M")
    s = _text(v)
    m = re.match(r"^(\d{1,2}):(\d{2})(?::\d{2})?$", s or "")
    if not m:
        raise ValueError(f"'{v}' is not a time (use HH:MM)")
    return f"{int(m.group(1)):02d}:{m.group(2)}"


def _split(v) -> list[str]:
    s = _text(v)
    return [x.strip() for x in re.split(r"[,;/\n]", s) if x.strip()] if s else []


def _pydantic_message(exc: ValidationError) -> str:
    return "; ".join(f"{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in exc.errors())


def commit(session: Session, store: ImportStore, token: str, mapping: dict, options: dict) -> dict:
    upload = store.get(token)
    for name, required, _d, _s in FIELDS[upload.kind]:
        if required and not mapping.get(name):
            raise HTTPException(status_code=422, detail=f"Map a column to the required field '{name}'")
    unknown = set(mapping) - {f[0] for f in FIELDS[upload.kind]}
    if unknown:
        raise HTTPException(status_code=422, detail=f"Unknown fields in mapping: {sorted(unknown)}")
    handler = {"analysts": _commit_analysts, "leave": _commit_leave, "roster": _commit_roster}[upload.kind]
    result = handler(session, upload, mapping, options)
    if not result["errors"]:
        store.drop(token)
    return result


def _rejected(errors: RowErrors) -> dict:
    return {"created": 0, "updated": 0, "skipped": 0, "errors": errors.items, "roster_id": None,
            "message": f"Nothing imported: {len(errors.items)} problem(s) found. Fix the file and try again."}


def _commit_analysts(session: Session, upload: Upload, mapping: dict, options: dict) -> dict:
    errors = RowErrors()
    projects = {p.code.upper(): p for p in session.scalars(select(Project))}
    skills = {s.name.lower(): s for s in session.scalars(select(Skill))}
    existing = {a.employee_code: a for a in session.scalars(select(Analyst))}
    update = bool(options.get("update_existing"))
    plans, seen = [], set()
    for rownum, rec, _row, _cols in _records(upload, mapping):
        try:
            code = _text(rec.get("employee_code"))
            proj_codes = [c.upper() for c in _split(rec.get("projects"))]
            primary = (_text(rec.get("primary_project")) or "").upper() or (proj_codes[0] if proj_codes else None)
            if primary and primary not in proj_codes:
                proj_codes.append(primary)
            missing = [c for c in proj_codes if c not in projects]
            if missing:
                raise ValueError(f"unknown project code(s): {', '.join(missing)}")
            active_txt = (_text(rec.get("active")) or "yes").lower()
            if active_txt not in ("yes", "y", "true", "1", "active", "no", "n", "false", "0", "inactive"):
                raise ValueError(f"active must be yes/no, got '{active_txt}'")
            hours = _text(rec.get("max_hours_per_week"))
            body = AnalystIn(
                employee_code=code or "", full_name=_text(rec.get("full_name")) or "",
                email=_text(rec.get("email")), active=active_txt in ("yes", "y", "true", "1", "active"),
                max_hours_per_week=int(hours) if hours else None,
                eligibility=[{"project_id": projects[c].id, "is_primary": c == primary} for c in proj_codes],
            )
            if body.employee_code in seen:
                raise ValueError(f"employee code {body.employee_code} appears twice in the file")
            seen.add(body.employee_code)
            skill_names = _split(rec.get("skills"))
            for s in skill_names:
                if len(s) > 64:
                    raise ValueError("skill names must be at most 64 characters")
            plans.append((body, skill_names))
        except ValidationError as exc:
            errors.add(rownum, _pydantic_message(exc))
        except ValueError as exc:
            errors.add(rownum, str(exc))
    if errors:
        return _rejected(errors)
    created = updated = skipped = 0
    for body, skill_names in plans:
        objs = []
        for s in skill_names:
            if s.lower() not in skills:
                skills[s.lower()] = Skill(name=s, description="Created by import")
                session.add(skills[s.lower()])
            objs.append(skills[s.lower()])
        current = existing.get(body.employee_code)
        if current is not None and not update:
            skipped += 1
            continue
        target = current or Analyst(employee_code=body.employee_code, preferences={})
        if current is None:
            session.add(target)
        target.full_name, target.email, target.active = body.full_name, body.email, body.active
        target.max_hours_per_week = body.max_hours_per_week
        target.skills = objs if (objs or current is None) else target.skills
        session.flush()
        if body.eligibility:
            target.eligibility.clear()
            session.flush()
            for e in body.eligibility:
                target.eligibility.append(AnalystProject(project_id=e.project_id, is_primary=e.is_primary,
                                                         allowed_shift_codes=[]))
        created += current is None
        updated += current is not None
    record(session, "import", "analysts", "", created=created, updated=updated)
    session.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "errors": [], "roster_id": None,
            "message": f"Imported analysts: {created} created, {updated} updated, {skipped} skipped (already exist)"}


def _commit_leave(session: Session, upload: Upload, mapping: dict, options: dict) -> dict:
    errors = RowErrors()
    analysts = {a.employee_code.upper(): a for a in session.scalars(select(Analyst))}
    plans = []
    for rownum, rec, _row, _cols in _records(upload, mapping):
        try:
            code = (_text(rec.get("employee_code")) or "").upper()
            if code not in analysts:
                raise ValueError(f"unknown employee code '{code}'")
            start = _parse_date(rec.get("start_date"))
            end = _parse_date(rec["end_date"]) if rec.get("end_date") is not None else start
            st, et = _parse_time(rec.get("start_time")), _parse_time(rec.get("end_time"))
            status = (_text(rec.get("status")) or "pending").lower()
            if status not in ("approved", "pending"):
                raise ValueError("status must be approved or pending")
            body = LeaveIn(analyst_id=analysts[code].id, start_date=start, end_date=end,
                           leave_type=(_text(rec.get("leave_type")) or "annual").lower(),
                           partial_day=bool(st or et), start_time=st, end_time=et, reason=_text(rec.get("reason")) or "")
            plans.append((rownum, body, status))
        except ValidationError as exc:
            errors.add(rownum, _pydantic_message(exc))
        except ValueError as exc:
            errors.add(rownum, str(exc))
    # Overlaps with stored leave or with other rows of the same file.
    active = list(session.scalars(select(LeaveRequest).where(LeaveRequest.status.in_(("pending", "approved")),
                                                             LeaveRequest.partial_day.is_(False))))
    accepted: list[LeaveIn] = []
    skipped = 0
    final = []
    for rownum, body, status in plans:
        if body.partial_day:
            final.append((body, status))
            continue
        dup = [lv for lv in active if lv.analyst_id == body.analyst_id and lv.start_date <= body.end_date
               and lv.end_date >= body.start_date]
        if dup:
            if all((lv.start_date, lv.end_date) == (body.start_date, body.end_date) for lv in dup):
                skipped += 1  # the same leave imported again
                continue
            errors.add(rownum, f"overlaps existing leave #{dup[0].id}")
            continue
        if any(o.analyst_id == body.analyst_id and o.start_date <= body.end_date and o.end_date >= body.start_date
               for o in accepted):
            errors.add(rownum, "overlaps another row in this file")
            continue
        accepted.append(body)
        final.append((body, status))
    if errors:
        return _rejected(errors)
    for body, status in final:
        session.add(LeaveRequest(**body.model_dump(), status=status, source="import",
                                 decided_at=datetime.now() if status == "approved" else None,
                                 decision_note="Imported as approved" if status == "approved" else ""))
    record(session, "import", "leave", "", created=len(final))
    session.commit()
    return {"created": len(final), "updated": 0, "skipped": skipped, "errors": [], "roster_id": None,
            "message": f"Imported {len(final)} leave record(s); {skipped} duplicate(s) skipped"}


def _commit_roster(session: Session, upload: Upload, mapping: dict, options: dict) -> dict:
    errors = RowErrors()
    try:
        year, month = int(options.get("year") or 0), int(options.get("month") or 0)
        start, end = month_bounds(year, month)
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="options.year and options.month are required") from None
    status = options.get("status") or "historical"
    if status not in ("historical", "draft"):
        raise HTTPException(status_code=422, detail="status must be historical or draft")
    if not mapping.get("employee_code") and not mapping.get("full_name"):
        raise HTTPException(status_code=422, detail="Map the analyst code or name column")
    default_project = session.get(Project, int(options["project_id"])) if options.get("project_id") else None
    projects = {p.code.upper(): p for p in session.scalars(select(Project))}
    templates = {(t.project_id, t.code.upper()): t for t in session.scalars(select(ShiftTemplate))}
    mappings = {m.source_code.upper(): m for m in session.scalars(select(ShiftCodeMapping))}
    by_code = {a.employee_code.upper(): a for a in session.scalars(select(Analyst))}
    by_name: dict[str, list[Analyst]] = {}
    for a in by_code.values():
        by_name.setdefault(a.full_name.strip().lower(), []).append(a)
    create_missing = bool(options.get("create_missing_analysts"))

    sheet = next(s for s in upload.sheets if s.name == upload.sheet)
    h = find_header(sheet.rows, "roster")
    hdr = sheet.rows[h]
    day_cols: dict[int, date] = {}
    for i, v in enumerate(hdr):
        if v is None or not _is_day_header(v):
            continue
        d = _date_header(v, year, month)
        if d is None:
            continue
        if not start <= d <= end:
            errors.add(h + 1, f"column '{v}' is {d}, outside {year}-{month:02d}")
            continue
        if d in day_cols.values():
            errors.add(h + 1, f"day {d} appears in two columns")
            continue
        day_cols[i] = d
    if not day_cols:
        errors.add(h + 1, "no day columns for the selected month")

    planned_analysts: dict[str, dict] = {}
    cells: list[tuple[str, date, str, object]] = []  # (analyst key, date, kind, payload)
    for rownum, rec, row, _cols in _records(upload, mapping):
        code = (_text(rec.get("employee_code")) or "").upper() if mapping.get("employee_code") else ""
        name = _text(rec.get("full_name")) if mapping.get("full_name") else None
        if not code and not name:
            continue
        if all((row[i] if i < len(row) else None) is None for i in day_cols):
            continue  # notes or totals below the grid, not an analyst row
        analyst = by_code.get(code) if code else None
        if analyst is None and name and not code:
            matches = by_name.get(name.lower(), [])
            if len(matches) > 1:
                errors.add(rownum, f"name '{name}' matches several analysts; add an employee code column")
                continue
            analyst = matches[0] if matches else None
        key = code or (name or "").lower()
        if analyst is None:
            if not create_missing:
                errors.add(rownum, f"analyst '{code or name}' not found (tick 'create missing analysts' to add)")
                continue
            if default_project is None:
                errors.add(rownum, "choose a project so new analysts can be made eligible")
                continue
            planned_analysts[key] = {"code": code or None, "name": name or code}
        else:
            planned_analysts[key] = {"existing": analyst}
        for i, d in day_cols.items():
            raw = row[i] if i < len(row) else None
            if raw is None:
                continue
            if isinstance(raw, FormulaCell):
                errors.add(rownum, f"{d}: formulas are not accepted")
                continue
            for token in _split(raw):
                proj = default_project
                code_part = token
                if ":" in token:
                    pcode, code_part = token.split(":", 1)
                    proj = projects.get(pcode.strip().upper())
                    if proj is None:
                        errors.add(rownum, f"{d}: unknown project '{pcode}' in '{token}'")
                        continue
                m = mappings.get(code_part.strip().upper())
                if m is None:
                    errors.add(rownum, f"{d}: unknown code '{code_part}' (add it under Settings > Code mappings)")
                    continue
                if m.target_kind != "shift":
                    cells.append((key, d, "marker", m.target_kind))
                    continue
                if proj is None:
                    errors.add(rownum, f"{d}: choose a project for shift code '{code_part}'")
                    continue
                tpl = templates.get((proj.id, (m.target_shift_code or "").upper()))
                if tpl is None:
                    errors.add(rownum, f"{d}: project {proj.code} has no shift template '{m.target_shift_code}'")
                    continue
                cells.append((key, d, "shift", tpl))
    seen_cells = set()
    for key, d, kind, payload in cells:
        ident = (key, d, kind, payload.id if kind == "shift" else payload)
        if ident in seen_cells:
            errors.add(None, f"{key} {d}: the same code appears twice")
        seen_cells.add(ident)
    for key, d in {(k, d) for k, d, kind, _ in cells if kind == "marker"}:
        if any(k == key and dd == d and kind == "shift" for k, dd, kind, _ in cells):
            errors.add(None, f"{key} {d}: a shift and a day-off code in the same cell")
    if errors:
        return _rejected(errors)

    analyst_ids: dict[str, int] = {}
    created_analysts = 0
    for key, info in planned_analysts.items():
        if "existing" in info:
            analyst_ids[key] = info["existing"].id
            continue
        n = created_analysts + 1
        code = info["code"] or f"IMP-{year}{month:02d}-{n:03d}"
        while code in by_code:
            n += 1
            code = f"IMP-{year}{month:02d}-{n:03d}"
        a = Analyst(employee_code=code, full_name=info["name"][:120], preferences={},
                    notes=f"Created by roster import ({upload.filename[:80]})")
        session.add(a)
        session.flush()
        session.add(AnalystProject(analyst_id=a.id, project_id=default_project.id, is_primary=True,
                                   allowed_shift_codes=[]))
        by_code[code] = a
        analyst_ids[key] = a.id
        created_analysts += 1
    project_ids = sorted({p.project_id for _k, _d, kind, p in cells if kind == "shift"}
                         | ({default_project.id} if default_project else set()))
    version = next_version(session, year, month)
    roster = Roster(year=year, month=month, status=status, version=version,
                    name=f"{year}-{month:02d} v{version} (imported)",
                    project_ids=project_ids, analyst_ids=sorted(set(analyst_ids.values())),
                    notes=f"Imported from {upload.filename[:100]}", solver_status=None)
    session.add(roster)
    session.flush()
    n_shifts = n_markers = 0
    for key, d, kind, payload in cells:
        if kind == "shift":
            session.add(Assignment(roster_id=roster.id, analyst_id=analyst_ids[key], project_id=payload.project_id,
                                   shift_template_id=payload.id, date=d, source="import"))
            n_shifts += 1
        else:
            session.add(RosterDayMarker(roster_id=roster.id, analyst_id=analyst_ids[key], date=d, code=payload))
            n_markers += 1
    record(session, "import", "roster", roster.id, assignments=n_shifts, markers=n_markers,
           analysts_created=created_analysts)
    session.commit()
    return {"created": n_shifts, "updated": 0, "skipped": 0, "errors": [], "roster_id": roster.id,
            "message": f"Imported roster with {n_shifts} shifts and {n_markers} day markers"
                       + (f"; created {created_analysts} analyst(s)" if created_analysts else "")}

