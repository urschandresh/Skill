"""Generate the SYNTHETIC sample files in sample-data/.

These imitate the *layout* of the previous SOC roster (analysts in rows, one column per day,
codes M/A/N/WO/L) with invented people and invented assignments. They are NOT the real
August 2026 roster, which has not been supplied. Run from the repo root:

    python scripts/make_synthetic_fixtures.py
"""

from __future__ import annotations

import csv
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parent.parent / "sample-data"
NAMES = ["Synthetic Analyst A", "Synthetic Analyst B", "Synthetic Analyst C", "Synthetic Analyst D",
         "Synthetic Analyst E", "Synthetic Analyst F", "Synthetic Analyst G"]


def historical_roster(path: Path, year: int = 2026, month: int = 8) -> None:
    first = date(year, month, 1)
    days = [(first + timedelta(days=i)) for i in range(31) if (first + timedelta(days=i)).month == month]
    wb = Workbook()
    ws = wb.active
    ws.title = "SOC Roster"
    ws["A1"] = f"SYNTHETIC FIXTURE - SOC Shift Roster {first:%B %Y} (invented data, not a real roster)"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "Codes: M 06-15, A 14-23, N 22-07, WO weekly off, L leave"
    ws.append([])
    ws.append(["S.No", "Emp Code", "Name", *[d.day for d in days]])
    # A simple rotating pattern: 5 days on the same shift, 2 days off, shifted per analyst.
    cycle = ["M"] * 5 + ["WO"] * 2 + ["A"] * 5 + ["WO"] * 2 + ["N"] * 5 + ["WO"] * 2
    leave_for = {2: range(10, 13), 5: range(20, 22)}  # analyst index -> leave days of month
    for i, name in enumerate(NAMES):
        row = [i + 1, f"SYN-H{i + 1:02d}", name]
        for d in days:
            code = cycle[(d.day - 1 + i * 3) % len(cycle)]
            if d.day in leave_for.get(i, ()):
                code = "L"
            row.append(code)
        ws.append(row)
    ws.cell(row=ws.max_row + 2, column=1, value="Coverage varies by day in this synthetic sample.")
    wb.save(path)


def analysts_csv(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Employee Code", "Full Name", "Email", "Projects", "Primary Project", "Skills", "Active"])
        for i in range(1, 8):
            w.writerow([f"SYN-{i:03d}", f"Synthetic Analyst {i}", "", "SOC1", "SOC1", "SIEM", "yes"])


def leave_csv(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Employee Code", "Start Date", "End Date", "Leave Type", "Status", "Start Time", "End Time", "Reason"])
        w.writerow(["SYN-002", "2026-11-09", "2026-11-11", "annual", "approved", "", "", "Synthetic"])
        w.writerow(["SYN-004", "2026-11-18", "2026-11-18", "casual", "pending", "", "", "Synthetic"])
        w.writerow(["SYN-006", "2026-11-05", "2026-11-05", "sick", "approved", "09:00", "13:00", "Synthetic partial"])


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    historical_roster(OUT / "SYNTHETIC_historical_soc_roster_2026-08.xlsx")
    analysts_csv(OUT / "SYNTHETIC_analysts.csv")
    leave_csv(OUT / "SYNTHETIC_leave_2026-11.csv")
    print("Wrote synthetic fixtures to", OUT)
